from pathlib import Path
import pdfplumber
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Documentation" / "Campus_One_Project_Report.docx"
IMG = ROOT / "Documentation" / "Images"
BRIEF_DIR = ROOT / "Documentation" / "BriefPages"
BRIEF = Path(r"C:\Users\Admin\Downloads\ESOFT_admin_module-assessment_resource_14529_14529-1768899903848-L3-DIIT Final Project - Final Project (2) 3.pdf")

BRIEF_DIR.mkdir(parents=True, exist_ok=True)

def set_font(run, name="Times New Roman", size=12, bold=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(6)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    set_font(p.add_run(text))
    return p

def heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(5)
    return p

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_font(r, size=12, bold=True)
        r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(values):
            if row_index % 2:
                shade(cells[i], "EAF2F8")
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.15
            set_font(p.add_run(str(text)), size=12)
            set_cell_margins(cells[i])
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return table

def code_block(doc, code):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cant_split = OxmlElement("w:cantSplit")
    table.rows[0]._tr.get_or_add_trPr().append(cant_split)
    shade(cell, "F2F2F2")
    set_cell_margins(cell, 120, 160, 120, 160)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    for index, line in enumerate(code.strip().splitlines()):
        if index:
            p.add_run().add_break()
        run = p.add_run(line)
        set_font(run, "Times New Roman", 12)
    doc.add_paragraph()

def page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    set_font(run, size=10)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])

doc = Document()
for section in doc.sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.15
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name, size in (("Title", 20), ("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
    style = styles[name]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    if name == "Title":
        p_pr = style._element.get_or_add_pPr()
        border = p_pr.find(qn("w:pBdr"))
        if border is not None:
            p_pr.remove(border)

# Original project brief first, as required.
p = doc.add_paragraph()
p.style = doc.styles["Title"]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("Original Project Brief"), size=20, bold=True)
body(doc, "The following pages reproduce the complete assessment brief supplied for the Campus One Digital Academy final project.")
doc.add_page_break()
with pdfplumber.open(BRIEF) as pdf:
    for index, page in enumerate(pdf.pages, start=1):
        image_path = BRIEF_DIR / f"brief-page-{index:02d}.png"
        page.to_image(resolution=130).save(image_path, format="PNG")
        doc.add_picture(str(image_path), width=Inches(6.05))
        cap = doc.add_paragraph(f"Original project brief page {index}")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(0)
        set_font(cap.add_run(), size=8)
        if index != len(pdf.pages):
            doc.add_page_break()

doc.add_section(WD_SECTION.NEW_PAGE)
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("Campus One Digital Academy Student Registration Management System"), size=20, bold=True)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("Final Project Report"), size=16, bold=True)
doc.add_picture(str(ROOT / "WinFormsApp1" / "Assets" / "CampusOneLogo.png"), width=Inches(2.1))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
for label in ("Student Name", "Student Registration Number", "Programme", "Submission Date"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run(f"{label}: ______________________________"), bold=(label in ("Student Name", "Student Registration Number")))
doc.add_page_break()

heading(doc, "Contents")
for item in ("1 Project Overview", "2 Requirements and Technology", "3 System Design", "4 User Interface Implementation",
             "5 Database Design", "6 Core Functionality", "7 Additional Features", "8 Validation and Error Handling",
             "9 Testing", "10 User Guide", "11 Problems and Solutions", "12 Conclusion", "13 Code Evidence"):
    body(doc, item)
doc.add_page_break()

heading(doc, "1 Project Overview")
body(doc, "This project implements a Windows desktop system for Campus One Digital Academy. It authenticates an administrator and provides facilities to register, retrieve, update and delete student records stored in Microsoft SQL Server. The completed application also includes a searchable student-records screen, summary totals, selection-based editing and CSV export.")
heading(doc, "1.1 Project Objectives", 2)
for text in ("Provide controlled access through the required Login form.", "Store student and parent information in the required Student database.",
             "Implement reliable Create, Read, Update and Delete operations using parameterized SQL.", "Present a clear interface that follows the supplied form structure.",
             "Add useful record browsing, searching, reporting and validation features."):
    bullet(doc, text)

heading(doc, "2 Requirements and Technology")
add_table(doc, ["Area", "Implementation"], [
    ("Development environment", "Microsoft Visual Studio compatible solution"),
    ("Language", "C#"), ("Framework", ".NET 8 Windows Forms"),
    ("Database", "Microsoft SQL Server / SQL Server Express"),
    ("Database access", "ADO.NET with Microsoft.Data.SqlClient 7.0.2"),
    ("Database name", "Student"), ("Main table", "Registration")], [2.0, 4.4])

heading(doc, "3 System Design")
body(doc, "The solution uses three forms. The Login form is the entry point. Successful authentication opens the Registration form, which performs all compulsory student operations. The Student Records form is an additional screen opened from Registration and displays database records in a DataGridView.")
add_table(doc, ["Form", "Purpose"], [
    ("frmLogin", "Authenticates the administrator and controls application entry."),
    ("frmRegistration", "Registers, retrieves, updates and deletes student records."),
    ("frmStudentRecords", "Lists, searches, summarizes, exports and selects records for editing.")], [1.8, 4.6])

heading(doc, "4 User Interface Implementation")
heading(doc, "4.1 Login Form", 2)
body(doc, "The Login form contains the required PictureBox, Login GroupBox, three labels, two textboxes and three buttons. The password is masked. Clear resets both fields, Login validates the supplied administrator credentials, and Exit requests confirmation.")
doc.add_picture(str(IMG / "LoginForm.png"), width=Inches(5.2))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
body(doc, "Figure 1 Login form")

heading(doc, "4.2 Student Registration Form", 2)
body(doc, "The Registration form contains the four required GroupBoxes and every field specified in the assessment. The enlarged layout separates labels from input controls and keeps Logout, action buttons and Exit visible.")
doc.add_picture(str(IMG / "RegistrationForm.png"), height=Inches(7.2))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
body(doc, "Figure 2 Student Registration form")

heading(doc, "4.3 Student Records Form", 2)
body(doc, "The additional Student Records form displays registered students in a read-only DataGridView. It supports live filtering, refresh, CSV export and returning the selected registration number to the Registration form for editing.")
doc.add_picture(str(IMG / "StudentRecordsForm.png"), width=Inches(6.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
body(doc, "Figure 3 Student Records form")

heading(doc, "5 Database Design")
body(doc, "The database is named Student and contains the Registration table. The registration number is the primary key. The application uses the exact field names and SQL Server types supplied in the assessment brief.")
add_table(doc, ["Field", "SQL type", "Purpose"], [
    ("regNo", "INT PRIMARY KEY", "Unique student registration number"), ("firstName", "VARCHAR(50)", "Student first name"),
    ("lastName", "VARCHAR(50)", "Student last name"), ("dateOfBirth", "DATETIME", "Student date of birth"),
    ("gender", "VARCHAR(50)", "Selected gender"), ("address", "VARCHAR(50)", "Student address"),
    ("email", "VARCHAR(50)", "Student email address"), ("mobilePhone", "INT", "Student mobile number"),
    ("homePhone", "INT", "Student home number"), ("parentName", "VARCHAR(50)", "Parent or guardian name"),
    ("nic", "VARCHAR(50)", "National identity number"), ("contactNo", "INT", "Parent contact number")], [1.4, 1.6, 3.4])

heading(doc, "6 Core Functionality")
for title, text in (
    ("6.1 Login", "The application compares the entered username and password with Admin and Campusone@123. Correct credentials hide Login and open Registration. Incorrect credentials display an error."),
    ("6.2 Register", "The Register operation validates every required field, checks for a duplicate registration number and inserts the record with parameters."),
    ("6.3 Search and Load", "Registration numbers are loaded into the ComboBox. Selecting or entering an existing number retrieves the complete student record."),
    ("6.4 Update", "Update modifies the matching database row and reports whether a record was found."),
    ("6.5 Delete", "Delete requires a registration number and confirmation before removing the matching row."),
    ("6.6 Clear Logout and Exit", "Clear resets fields and focus. Logout returns to Login. Exit closes the application only after confirmation.")):
    heading(doc, title, 2)
    body(doc, text)

heading(doc, "7 Additional Features")
add_table(doc, ["Feature", "Benefit"], [
    ("Student Records DataGridView", "Provides a complete, sortable overview of stored students."),
    ("Live search", "Filters by registration number, first name, last name or NIC while typing."),
    ("Dashboard totals", "Shows total, male and female counts for the displayed records."),
    ("Edit selected record", "Returns the selected registration number to the main form."),
    ("CSV export", "Creates a portable spreadsheet-compatible copy of displayed records."),
    ("Window branding", "Uses the supplied Campus One logo and title-bar icon.")], [2.2, 4.2])

heading(doc, "8 Validation and Error Handling")
body(doc, "The Registration form uses ErrorProvider messages beside invalid controls. It validates positive registration numbers, required names and addresses, email format, numeric phone fields, gender selection, NIC entry and future dates. SQL operations use exception handling and display clear database messages.")

heading(doc, "9 Testing")
add_table(doc, ["Test", "Expected result", "Status"], [
    ("Start application", "Login form opens first", "Pass"),
    ("Correct login", "Registration form opens", "Pass"),
    ("Incorrect login", "Login error appears", "Pass"),
    ("Clear login", "Credentials clear and Username receives focus", "Pass"),
    ("Register valid student", "Record is inserted and success appears", "Pass"),
    ("Duplicate Reg No", "Duplicate warning appears", "Pass"),
    ("Select Reg No", "Student details populate", "Pass"),
    ("Update record", "Matching row changes", "Pass"),
    ("Delete No", "Record remains", "Pass"),
    ("Delete Yes", "Record is removed", "Pass"),
    ("Live search", "Grid filters while typing", "Pass"),
    ("CSV export", "Displayed rows are written to CSV", "Pass"),
    ("Release build", "Build completes with zero code warnings or errors", "Pass")], [1.7, 3.5, 1.2])

heading(doc, "10 User Guide")
for text in ("Start SQL Server Express and run DatabaseScript.sql in SQL Server Management Studio if the Student database has not been created.",
             "Run the application and log in with Username Admin and Password Campusone@123.",
             "Enter all student, contact and parent information, then select Register.",
             "Choose a registration number to retrieve a student. Change the details and select Update when required.",
             "Select Delete and confirm Yes to remove the selected student.",
             "Select View Students to search, export or choose a record for editing.",
             "Use Logout to return to Login and Exit to close the program safely."):
    bullet(doc, text)

heading(doc, "11 Problems and Solutions")
add_table(doc, ["Problem", "Solution"], [
    ("Registration table not found", "Added an idempotent database script and startup schema check."),
    ("Labels overlapped input fields", "Enlarged the form and separated label and textbox coordinates."),
    ("Logout overlapped the logo", "Kept branding in the title bar and repositioned Logout."),
    ("Limited record navigation", "Added a searchable DataGridView screen with selection-based editing."),
    ("Generic validation feedback", "Added field-specific ErrorProvider messages and format checks.")], [2.4, 4.0])

doc.add_page_break()
heading(doc, "12 Conclusion")
body(doc, "The completed system satisfies the compulsory Login and Student Registration requirements and connects the required workflow to Microsoft SQL Server through parameterized ADO.NET commands. The additional records screen, search, summary, CSV export and detailed validation improve usability while preserving the required database structure.")

heading(doc, "13 Code Evidence")
heading(doc, "13.1 Login Form Authentication", 2)
code_block(doc, '''if (username == AdminUsername && password == AdminPassword)
{
    Hide();
    frmRegistration registrationForm = new frmRegistration();
    registrationForm.FormClosed += (s, args) => Close();
    registrationForm.Show();
}
else
{
    MessageBox.Show("Invalid Username or Password. Please try again.");
}''')

heading(doc, "13.2 Registration Insert", 2)
code_block(doc, '''string sql = @"INSERT INTO Registration
(regNo, firstName, lastName, dateOfBirth, gender, address, email,
 mobilePhone, homePhone, parentName, nic, contactNo)
VALUES
(@regNo, @firstName, @lastName, @dateOfBirth, @gender, @address,
 @email, @mobilePhone, @homePhone, @parentName, @nic, @contactNo)";
using SqlCommand command = new SqlCommand(sql, connection);
AddStudentParameters(command);
command.ExecuteNonQuery();''')

heading(doc, "13.3 Student Records Search", 2)
code_block(doc, '''SELECT regNo, firstName, lastName, dateOfBirth, gender, address,
       email, mobilePhone, homePhone, parentName, nic, contactNo
FROM Registration
WHERE @search = '' OR CONVERT(VARCHAR(20), regNo) LIKE @pattern
   OR firstName LIKE @pattern OR lastName LIKE @pattern OR nic LIKE @pattern
ORDER BY regNo''')

heading(doc, "13.4 Database Creation", 2)
code_block(doc, (ROOT / "WinFormsApp1" / "DatabaseScript.sql").read_text(encoding="utf-8"))

page_number(doc.sections[0].footer.paragraphs[0])

doc.core_properties.title = "Campus One Digital Academy Student Registration Management System"
doc.core_properties.subject = "Final Project Report"
doc.core_properties.author = "Student"
doc.core_properties.keywords = "C#, Windows Forms, SQL Server, Student Registration"
doc.save(OUT)
print(OUT)
